from __future__ import annotations

import json
import os
import posixpath
import re
import shutil
import sys
import zipfile
from datetime import date, datetime, time, timezone
from pathlib import Path
from typing import Any
from xml.etree import ElementTree as ET

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SETTINGS_FILE = ROOT / "settings.json"
DATA_DIR = ROOT / "data"
IMAGES_DIR = DATA_DIR / "images"
OUTPUT_FILE = DATA_DIR / "data.json"


def load_settings() -> dict[str, Any]:
    if not SETTINGS_FILE.exists():
        raise FileNotFoundError(
            f"Missing {SETTINGS_FILE.name}. Copy settings.example.json to settings.json "
            "and edit the Excel path and sheet name."
        )
    with SETTINGS_FILE.open("r", encoding="utf-8") as f:
        return json.load(f)


def json_safe(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, (int, float, bool, str)):
        return value
    return str(value)


def safe_name(value: Any, fallback: str) -> str:
    text = str(value or "").strip()
    text = re.sub(r'[<>:"/\\|?*\x00-\x1F]+', "_", text)
    text = re.sub(r"\s+", "_", text).strip("._ ")
    return (text[:80] or fallback)


def local_name(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def first_existing(names: list[str], candidates: list[str]) -> str | None:
    name_set = set(names)
    for candidate in candidates:
        if candidate in name_set:
            return candidate
    return None


def resolve_target(base_file: str, target: str) -> str:
    base_dir = posixpath.dirname(base_file)
    return posixpath.normpath(posixpath.join(base_dir, target))


def workbook_sheet_xml(zf: zipfile.ZipFile, sheet_name: str) -> str | None:
    names = zf.namelist()
    if "xl/workbook.xml" not in names or "xl/_rels/workbook.xml.rels" not in names:
        return None

    wb_root = ET.fromstring(zf.read("xl/workbook.xml"))
    rel_root = ET.fromstring(zf.read("xl/_rels/workbook.xml.rels"))

    rel_map = {}
    for rel in rel_root:
        rid = rel.attrib.get("Id")
        target = rel.attrib.get("Target")
        if rid and target:
            rel_map[rid] = target

    rid_attr = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"

    for el in wb_root.iter():
        if local_name(el.tag) == "sheet" and el.attrib.get("name") == sheet_name:
            rid = el.attrib.get(rid_attr)
            target = rel_map.get(rid)
            if not target:
                return None
            if target.startswith("/"):
                return target.lstrip("/")
            return posixpath.normpath(posixpath.join("xl", target))
    return None


def extract_modern_in_cell_images(
    excel_path: Path,
    sheet_name: str,
    primary_by_row: dict[int, Any],
) -> tuple[dict[int, list[str]], set[str], list[str]]:
    """
    Best-effort extractor for modern Excel "Place in Cell" images.

    It follows the Rich Data chain:
    worksheet c/@vm -> metadata.xml -> rich value -> richValueRel -> xl/media.

    Excel has more than one on-disk Rich Data variant. This function supports
    common rdRichValue/richValue variants and fails gracefully when a workbook
    uses an unsupported layout.
    """
    by_row: dict[int, list[str]] = {}
    image_cells: set[str] = set()
    warnings: list[str] = []

    try:
        with zipfile.ZipFile(excel_path, "r") as zf:
            names = zf.namelist()
            sheet_xml = workbook_sheet_xml(zf, sheet_name)
            metadata_xml = first_existing(names, ["xl/metadata.xml"])
            rv_struct_xml = first_existing(names, [
                "xl/richData/rdrichvaluestructure.xml",
                "xl/richData/richValueStructure.xml",
                "xl/richData/richvaluestructure.xml",
            ])
            rv_data_xml = first_existing(names, [
                "xl/richData/rdrichvalue.xml",
                "xl/richData/richValue.xml",
                "xl/richData/richvalue.xml",
            ])
            rv_rel_xml = first_existing(names, [
                "xl/richData/richValueRel.xml",
                "xl/richData/richvaluerel.xml",
            ])
            rv_rels_xml = first_existing(names, [
                "xl/richData/_rels/richValueRel.xml.rels",
                "xl/richData/_rels/richvaluerel.xml.rels",
            ])

            required = [sheet_xml, metadata_xml, rv_struct_xml, rv_data_xml, rv_rel_xml, rv_rels_xml]
            if not all(required):
                return by_row, image_cells, warnings

            # 1) cells carrying value-metadata (vm)
            sheet_root = ET.fromstring(zf.read(sheet_xml))
            cell_vm: list[tuple[str, int]] = []
            for c in sheet_root.iter():
                if local_name(c.tag) != "c":
                    continue
                ref = c.attrib.get("r")
                vm = c.attrib.get("vm")
                if ref and vm is not None:
                    try:
                        cell_vm.append((ref, int(vm)))
                    except ValueError:
                        pass

            if not cell_vm:
                return by_row, image_cells, warnings

            # Workbooks observed in the wild can use vm 0-based or 1-based.
            vm_base = 0 if any(vm == 0 for _, vm in cell_vm) else 1

            # 2) metadata: valueMetadata -> futureMetadata(XLRICHVALUE) -> rvb/@i
            meta_root = ET.fromstring(zf.read(metadata_xml))

            value_records = []
            future_rich_bks = []

            for el in meta_root.iter():
                if local_name(el.tag) == "valueMetadata":
                    value_records = [child for child in list(el) if local_name(child.tag) == "bk"]
                    break

            for el in meta_root.iter():
                if local_name(el.tag) == "futureMetadata" and el.attrib.get("name") == "XLRICHVALUE":
                    future_rich_bks = [child for child in list(el) if local_name(child.tag) == "bk"]
                    break

            def rich_index_from_vm(vm: int) -> int | None:
                idx = vm - vm_base
                if idx < 0 or idx >= len(value_records):
                    # fallback for alternate indexing
                    candidates = [vm, vm - 1]
                    idx = next((x for x in candidates if 0 <= x < len(value_records)), -1)
                if idx < 0:
                    return None

                rc = next((x for x in value_records[idx].iter() if local_name(x.tag) == "rc"), None)
                if rc is None:
                    return None
                try:
                    future_index_raw = int(rc.attrib.get("v", "0"))
                except ValueError:
                    return None

                candidates = [future_index_raw]
                if future_index_raw > 0:
                    candidates.append(future_index_raw - 1)

                future_bk = None
                for candidate in candidates:
                    if 0 <= candidate < len(future_rich_bks):
                        future_bk = future_rich_bks[candidate]
                        break
                if future_bk is None:
                    return None

                rvb = next((x for x in future_bk.iter() if local_name(x.tag) == "rvb"), None)
                if rvb is None:
                    return None
                try:
                    return int(rvb.attrib["i"])
                except (KeyError, ValueError):
                    return None

            # 3) rich value structures
            struct_root = ET.fromstring(zf.read(rv_struct_xml))
            structures = [el for el in struct_root if local_name(el.tag) == "s"]
            if not structures:
                structures = [el for el in struct_root.iter() if local_name(el.tag) == "s"]

            # 4) rich values
            rv_root = ET.fromstring(zf.read(rv_data_xml))
            rich_values = [el for el in rv_root if local_name(el.tag) == "rv"]
            if not rich_values:
                rich_values = [el for el in rv_root.iter() if local_name(el.tag) == "rv"]

            # 5) ordered relationship slots
            rel_slot_root = ET.fromstring(zf.read(rv_rel_xml))
            rid_attr = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
            rel_slots = [
                el.attrib.get(rid_attr)
                for el in rel_slot_root.iter()
                if local_name(el.tag) == "rel" and el.attrib.get(rid_attr)
            ]

            # 6) relationship id -> media target
            rels_root = ET.fromstring(zf.read(rv_rels_xml))
            rid_to_target = {}
            for rel in rels_root:
                rid = rel.attrib.get("Id")
                target = rel.attrib.get("Target")
                if rid and target:
                    rid_to_target[rid] = resolve_target(rv_rel_xml, target)

            for cell_ref, vm in cell_vm:
                rich_idx = rich_index_from_vm(vm)
                if rich_idx is None or not (0 <= rich_idx < len(rich_values)):
                    continue

                rv = rich_values[rich_idx]
                try:
                    struct_idx = int(rv.attrib.get("s", "0"))
                except ValueError:
                    struct_idx = 0

                if not (0 <= struct_idx < len(structures)):
                    continue

                structure = structures[struct_idx]
                if structure.attrib.get("t") != "_localImage":
                    continue

                keys = [
                    k.attrib.get("n", "")
                    for k in structure
                    if local_name(k.tag) == "k"
                ]
                values = [
                    (v.text or "").strip()
                    for v in rv
                    if local_name(v.tag) == "v"
                ]

                try:
                    key_pos = next(
                        i for i, key in enumerate(keys)
                        if "LocalImageIdentifier" in key
                    )
                except StopIteration:
                    continue

                if key_pos >= len(values):
                    continue

                try:
                    slot_idx = int(values[key_pos])
                except ValueError:
                    continue

                if not (0 <= slot_idx < len(rel_slots)):
                    continue

                rid = rel_slots[slot_idx]
                media_path = rid_to_target.get(rid)
                if not media_path or media_path not in names:
                    continue

                match = re.search(r"(\d+)$", cell_ref)
                if not match:
                    continue
                row_num = int(match.group(1))

                ext = Path(media_path).suffix.lower() or ".png"
                primary = safe_name(primary_by_row.get(row_num), f"row_{row_num}")
                filename = f"{primary}_row{row_num}_cell_{cell_ref}{ext}"
                out_path = IMAGES_DIR / filename

                with out_path.open("wb") as f:
                    f.write(zf.read(media_path))

                relative = out_path.relative_to(ROOT).as_posix()
                by_row.setdefault(row_num, []).append(relative)
                image_cells.add(cell_ref)

    except Exception as exc:
        warnings.append(f"Modern in-cell image extraction warning: {exc}")

    return by_row, image_cells, warnings


def extract_openpyxl_images(ws, primary_by_row: dict[int, Any]) -> tuple[dict[int, list[str]], set[str], list[str]]:
    by_row: dict[int, list[str]] = {}
    image_cells: set[str] = set()
    warnings: list[str] = []

    images = getattr(ws, "_images", [])
    for number, img in enumerate(images, start=1):
        try:
            anchor = getattr(img, "anchor", None)
            marker = getattr(anchor, "_from", None)
            if marker is None:
                continue

            row_num = marker.row + 1
            col_num = marker.col + 1
            cell_ref = ws.cell(row_num, col_num).coordinate

            raw = img._data()
            ext = (getattr(img, "format", "") or "").lower()
            if ext == "jpeg":
                ext = "jpg"
            if ext not in {"png", "jpg", "gif", "bmp", "tiff", "webp"}:
                ext = "png"

            primary = safe_name(primary_by_row.get(row_num), f"row_{row_num}")
            filename = f"{primary}_row{row_num}_img{number}.{ext}"
            out_path = IMAGES_DIR / filename

            with out_path.open("wb") as f:
                f.write(raw)

            relative = out_path.relative_to(ROOT).as_posix()
            by_row.setdefault(row_num, []).append(relative)
            image_cells.add(cell_ref)
        except Exception as exc:
            warnings.append(f"Image {number}: {exc}")

    return by_row, image_cells, warnings


def main() -> None:
    settings = load_settings()
    excel_path = Path(os.path.expandvars(settings["excel_file"])).expanduser()
    sheet_name = settings.get("sheet_name", "Sheet1")
    header_row = int(settings.get("header_row", 1))

    if not excel_path.exists():
        raise FileNotFoundError(f"Excel file not found: {excel_path}")

    DATA_DIR.mkdir(parents=True, exist_ok=True)
    if IMAGES_DIR.exists():
        shutil.rmtree(IMAGES_DIR)
    IMAGES_DIR.mkdir(parents=True, exist_ok=True)

    # Keep formulas' cached values where possible.
    wb = load_workbook(excel_path, data_only=True)
    if sheet_name not in wb.sheetnames:
        raise KeyError(
            f"Sheet '{sheet_name}' not found. Available sheets: {', '.join(wb.sheetnames)}"
        )
    ws = wb[sheet_name]

    raw_headers = [ws.cell(header_row, col).value for col in range(1, ws.max_column + 1)]
    headers: list[str] = []
    used = set()
    for col_idx, raw in enumerate(raw_headers, start=1):
        base = str(raw).strip() if raw not in (None, "") else f"Column {col_idx}"
        name = base
        n = 2
        while name in used:
            name = f"{base} {n}"
            n += 1
        used.add(name)
        headers.append(name)

    # Use the first non-empty column as a friendly image filename prefix.
    first_data_col = 1
    for idx, header in enumerate(raw_headers, start=1):
        if header not in (None, ""):
            first_data_col = idx
            break

    primary_by_row = {
        row: ws.cell(row, first_data_col).value
        for row in range(header_row + 1, ws.max_row + 1)
    }

    legacy_images, legacy_cells, warnings1 = extract_openpyxl_images(ws, primary_by_row)
    modern_images, modern_cells, warnings2 = extract_modern_in_cell_images(
        excel_path, sheet_name, primary_by_row
    )

    images_by_row: dict[int, list[str]] = {}
    for source in (legacy_images, modern_images):
        for row_num, paths in source.items():
            for p in paths:
                if p not in images_by_row.setdefault(row_num, []):
                    images_by_row[row_num].append(p)

    image_cells = legacy_cells | modern_cells
    rows = []

    for row_num in range(header_row + 1, ws.max_row + 1):
        record = {}
        nonempty = False

        for col_idx, header in enumerate(headers, start=1):
            cell = ws.cell(row_num, col_idx)
            value = json_safe(cell.value)

            # Modern in-cell images can expose placeholder values such as 0/#VALUE!.
            if cell.coordinate in image_cells:
                value = ""

            if value not in ("", None):
                nonempty = True
            record[header] = value

        row_images = images_by_row.get(row_num, [])
        if row_images:
            nonempty = True
            record["__images"] = row_images
            record["__image"] = row_images[0]
        else:
            record["__images"] = []

        record["__row_number"] = row_num

        if nonempty:
            rows.append(record)

    payload = {
        "meta": {
            "generated_at": datetime.now(timezone.utc).isoformat(),
            "source_file": excel_path.name,
            "sheet": sheet_name,
            "record_count": len(rows),
        },
        "rows": rows,
    }

    with OUTPUT_FILE.open("w", encoding="utf-8") as f:
        json.dump(payload, f, ensure_ascii=False, indent=2)

    print(f"Export complete: {len(rows)} rows")
    print(f"Images exported: {sum(len(v) for v in images_by_row.values())}")
    print(f"Website data: {OUTPUT_FILE}")

    for warning in warnings1 + warnings2:
        print(f"WARNING: {warning}")


if __name__ == "__main__":
    try:
        main()
    except Exception as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        raise
